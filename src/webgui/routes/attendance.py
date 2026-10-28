"""
Attendance Routes - WebGUI
Rutas para visualización y exportación de registros de asistencia
"""

import csv
import io
from datetime import datetime
from typing import Any

from flask import (
    Blueprint,
    Response,
    current_app,
    jsonify,
    render_template,
    request,
    send_file,
)
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ...utils.timezones import coerce_timezone, get_system_timezone, to_local_iso
from ..legal_export import LEGAL_HEADERS, build_legal_rows
from ..middleware.auth import api_require, require

attendance_bp = Blueprint("attendance", __name__)


def localize_rows(rows: list[dict], db) -> None:
    """Convierte timestamps de rows a la hora local del centro de trabajo (MYK-420)."""
    tz = coerce_timezone(get_system_timezone(db))
    for row in rows:
        row["timestamp"] = to_local_iso(row["timestamp"], tz)


def get_buffer_repo():
    """Obtener BufferRepository desde current_app"""
    return current_app.buffer_repo


def get_webgui_db():
    """Obtener Database (webgui.db) desde current_app"""
    return current_app.db


def resolve_employee_names(db, rows) -> dict[tuple[str, str], str]:
    """
    Resuelve el nombre del colaborador vía employee_device_links.

    Mapea (device_ip, user_id) -> nombre apellido usando el enlace
    device_serial (via devices.address) + device_user_id. Los registros sin
    enlace quedan fuera del mapa (fallback a device_users en el caller).
    """
    pairs = {(row["device_ip"], row["user_id"]) for row in rows if row["device_ip"]}

    device_serial_by_ip: dict[str, str] = {}
    employee_by_link: dict[tuple[str, str], str] = {}

    if pairs:
        device_rows = db.fetch_all(
            "SELECT address, serial_number FROM devices WHERE serial_number IS NOT NULL"
        )
        device_serial_by_ip = {d["address"]: d["serial_number"] for d in device_rows}

        serials = {device_serial_by_ip[ip] for ip, _ in pairs if ip in device_serial_by_ip}
        if serials:
            placeholders = ",".join("?" * len(serials))
            links = db.fetch_all(
                f"""
                SELECT edl.device_serial, edl.device_user_id,
                       e.first_name, e.last_name
                FROM employee_device_links edl
                JOIN employees e ON edl.employee_id = e.id
                WHERE edl.device_serial IN ({placeholders})
                """,
                tuple(serials),
            )
            employee_by_link = {
                (
                    row["device_serial"],
                    row["device_user_id"],
                ): f"{row['first_name']} {row['last_name']}".strip()
                for row in links
            }

    result: dict[tuple[str, str], str] = {}
    for device_ip, user_id in pairs:
        serial = device_serial_by_ip.get(device_ip)
        if serial and (serial, user_id) in employee_by_link:
            result[(device_ip, user_id)] = employee_by_link[(serial, user_id)]
    return result


@attendance_bp.route("/", strict_slashes=False)
@require("reports:read")
def list_attendance():
    """Página principal de asistencia - renderiza template"""
    return render_template("attendance/list.html")


@attendance_bp.route("/api/devices")
@api_require("reports:read")
def api_devices():
    """API para obtener lista de dispositivos para el select de filtro"""
    db = get_webgui_db()
    # Usar columnas reales del schema: address, is_enabled (migracion 001)
    devices = db.fetch_all(
        "SELECT id, name, address AS ip, is_enabled AS enabled "
        "FROM devices WHERE is_enabled = 1 ORDER BY name"
    )
    return jsonify({"devices": devices})


@attendance_bp.route("/api/attendance")
@api_require("reports:read")
def api_attendance():
    """
    API paginada con filtros para registros de asistencia

    Query params:
    - page: página (default 1)
    - per_page: registros por página (default 50, max 200)
    - start_date: fecha inicio ISO (YYYY-MM-DD)
    - end_date: fecha fin ISO (YYYY-MM-DD)
    - user_id: filtro por user_id
    - device_ip: filtro por IP de dispositivo
    - status: filtro por status (0-5)
    - punch_type: filtro por tipo de marca
    - synced: filtro por sincronizado (0/1)
    - sort: campo para ordenar
    - order: asc/desc
    """
    buffer_repo = get_buffer_repo()
    db = get_webgui_db()

    # Parámetros de paginación
    page = max(1, request.args.get("page", 1, type=int))
    per_page = min(max(10, request.args.get("per_page", 50, type=int)), 200)
    offset = (page - 1) * per_page

    # Parámetros de filtro
    start_date = request.args.get("start_date", "")
    end_date = request.args.get("end_date", "")
    user_id = request.args.get("user_id", "")
    device_ip = request.args.get("device_ip", "")
    status = request.args.get("status", "")
    punch_type = request.args.get("punch_type", "")
    synced = request.args.get("synced", "")
    sort = request.args.get("sort", "timestamp")
    order = request.args.get("order", "desc")

    # Validar sort field
    allowed_sorts = ["id", "user_id", "timestamp", "status", "device_ip", "punch_type", "synced"]
    if sort not in allowed_sorts:
        sort = "timestamp"
    if order.lower() not in ["asc", "desc"]:
        order = "desc"

    # Construir WHERE clause. Los registros anulados por correccion (MYK-422)
    # no aparecen en listados/exportaciones: su historial vive en la bitacora.
    where_conditions = []
    params = []

    if start_date:
        where_conditions.append("substr(timestamp, 1, 10) >= ?")
        params.append(start_date)
    if end_date:
        where_conditions.append("substr(timestamp, 1, 10) <= ?")
        params.append(end_date)
    if user_id:
        where_conditions.append("user_id = ?")
        params.append(user_id)
    if device_ip:
        where_conditions.append("device_ip = ?")
        params.append(device_ip)
    if status != "":
        where_conditions.append("status = ?")
        params.append(int(status))
    if punch_type:
        where_conditions.append("punch_type = ?")
        params.append(punch_type)
    if synced != "":
        where_conditions.append("synced = ?")
        params.append(int(synced))

    # Los anulados por correccion no se listan (MYK-422). En buffers legados
    # sin migration 010 la columna no existe: modo degradado sin filtro.
    if _buffer_has_column(buffer_repo, "attendance_records", "annulled"):
        where_conditions.append("annulled = 0")

    where_clause = "WHERE " + " AND ".join(where_conditions) if where_conditions else ""

    # Contar total
    count_query = f"SELECT COUNT(*) FROM attendance_records {where_clause}"
    with buffer_repo.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(count_query, params)
        total = cursor.fetchone()[0]

    # Obtener registros paginados
    data_query = f"""
        SELECT id, user_id, timestamp, status, device_ip, punch_type, synced
        FROM attendance_records
        {where_clause}
        ORDER BY {sort} {order.upper()}
        LIMIT ? OFFSET ?
    """
    params_with_pagination = params + [per_page, offset]

    with buffer_repo.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(data_query, params_with_pagination)
        rows = cursor.fetchall()

    records = [dict(row) for row in rows]

    # Localize timestamps to system timezone (MYK-420)
    localize_rows(records, db)

    return jsonify(
        {
            "records": records,
            "pagination": {
                "page": page,
                "per_page": per_page,
                "total": total,
                "pages": (total + per_page - 1) // per_page,
            },
        }
    )


def _buffer_has_column(repo: Any, table: str, column: str) -> bool:
    with repo.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(f"PRAGMA table_info({table})")
        return any(row[1] == column for row in cursor.fetchall())


def _parse_export_date(value: str | None, field: str) -> str | None:
    if not value:
        return None
    try:
        datetime.strptime(value.strip(), "%Y-%m-%d")
    except ValueError as exc:
        msg = f"{field} invalida: {value!r} (formato esperado YYYY-MM-DD)"
        raise ExportValidationError(msg) from exc
    return value.strip()


class ExportValidationError(ValueError):
    """Parametros invalidos para el export legal."""


EXPORT_FORMATS: dict[str, str] = {"csv": "csv", "xlsx": "xlsx", "excel": "xlsx"}


@attendance_bp.route("/export")
@api_require("reports:export")
def export():
    """Exportar asistencia a CSV o Excel (formato legal, MYK-424).

    Columnas: codigo de empleado, colaborador, CURP, RFC, razon social,
    RFC patron, centro de trabajo (ubicacion del dispositivo, modo degradado
    hasta CT-1/MYK-427), fecha-hora con zona (MYK-420), tipo de marca, estado
    y desglose de jornada semanal LFT 66-68 (ordinaria / extra doble / extra
    triple). Sin registros devuelve 200 con solo la cabecera (CSV/XLSX validos).
    """
    repo = get_buffer_repo()
    db = get_webgui_db()

    try:
        start_date = _parse_export_date(request.args.get("start_date"), "start_date")
        end_date = _parse_export_date(request.args.get("end_date"), "end_date")
    except ExportValidationError as exc:
        return str(exc), 400
    user_id = request.args.get("user_id") or ""
    device_ip = request.args.get("device_ip") or ""
    status = request.args.get("status") or ""
    punch_type = request.args.get("punch_type") or ""
    synced = request.args.get("synced") or ""
    fmt = EXPORT_FORMATS.get((request.args.get("format") or "csv").lower())
    if fmt is None:
        return "Formato de exportacion no soportado: csv|xlsx", 400

    where_conditions: list[str] = []
    params: list[Any] = []
    if start_date:
        where_conditions.append("substr(timestamp, 1, 10) >= ?")
        params.append(start_date)
    if end_date:
        where_conditions.append("substr(timestamp, 1, 10) <= ?")  # end_date inclusivo (MYK-383)
        params.append(end_date)
    if user_id:
        where_conditions.append("user_id = ?")
        params.append(user_id)
    if device_ip:
        where_conditions.append("device_ip = ?")
        params.append(device_ip)
    if status != "":
        try:
            status_int = int(status)
        except ValueError:
            return "Estado invalido: debe ser 0-5", 400
        where_conditions.append("status = ?")
        params.append(status_int)
    if punch_type:
        where_conditions.append("punch_type = ?")
        params.append(punch_type)
    if synced != "":
        try:
            params.append(int(synced))
        except ValueError:
            return "Sincronizado invalido: use 0 o 1", 400
        where_conditions.append("synced = ?")

    # Los registros anulados por correccion (MYK-422) no integran el export
    # legal: su historial vive en la bitacora. La columna solo existe en
    # buffers migrados con migration 010 (restaurada por MYK-479); en buffers
    # legados se exportan todos (modo degradado, sin romper el endpoint).
    if _buffer_has_column(repo, "attendance_records", "annulled"):
        where_conditions.append("annulled = 0")

    where_clause = "WHERE " + " AND ".join(where_conditions) if where_conditions else ""
    data_query = f"""
        SELECT id, user_id, timestamp, status, device_ip, punch_type, synced
        FROM attendance_records
        {where_clause}
        ORDER BY timestamp ASC, id ASC
    """
    with repo.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(data_query, params)
        records = [dict(row) for row in cursor.fetchall()]

    tz = coerce_timezone(get_system_timezone(db))
    rows = build_legal_rows(records, db, repo, tz)

    start_label = start_date or "todos"
    end_label = end_date or "todos"
    download_name = f"asistencia_legal_{start_label}_{end_label}.{fmt}"

    if fmt == "xlsx":
        return _render_xlsx(rows, download_name)
    return _render_csv(rows, download_name)


def _render_csv(rows: list[dict[str, str]], download_name: str) -> Response:
    output = io.StringIO()
    output.write("\ufeff")  # BOM utf-8-sig para Excel
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(LEGAL_HEADERS)
    for row in rows:
        writer.writerow([row.get(header, "") for header in LEGAL_HEADERS])

    response = Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{download_name}"'},
    )
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    return response


def _render_xlsx(rows: list[dict[str, str]], download_name: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia legal"

    ws.append(LEGAL_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in LEGAL_HEADERS])

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_length + 2, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=download_name,
    )
