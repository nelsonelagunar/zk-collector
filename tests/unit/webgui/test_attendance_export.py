"""
Unit tests for Attendance Export routes
"""

import csv
import io
from datetime import datetime, timedelta

import pytest
from flask import Flask
from openpyxl import load_workbook

from src.buffer.models import AttendanceRecordDB
from src.buffer.repository import BufferRepository


@pytest.fixture
def app(tmp_path):
    """Create a Flask app for testing"""
    db_path = tmp_path / "webgui.db"
    buffer_db_path = tmp_path / "buffer.db"

    app = Flask(__name__)
    app.config["SECRET_KEY"] = "test-secret-key"
    app.config["TESTING"] = True

    # Mock buffer repository (empty by default)
    buffer_repo = BufferRepository(buffer_db_path)
    app.buffer_repo = buffer_repo

    # Register attendance blueprint
    from src.webgui.routes.attendance import attendance_bp

    app.register_blueprint(attendance_bp)

    return app


@pytest.fixture
def app_with_data(app, tmp_path, sample_records):
    """Create a Flask app with populated buffer repository"""
    buffer_db_path = tmp_path / "buffer_test.db"
    repo = BufferRepository(buffer_db_path)
    repo.add_records(sample_records)
    app.buffer_repo = repo
    return app


@pytest.fixture
def client(app):
    """Create a test client"""
    return app.test_client()


@pytest.fixture
def sample_records():
    """Create sample attendance records for testing"""
    now = datetime.now()
    return [
        AttendanceRecordDB(
            user_id="1",
            timestamp=now - timedelta(days=2, hours=8),
            status=0,
            device_ip="192.168.1.101",
            punch_type="entrada",
        ),
        AttendanceRecordDB(
            user_id="2",
            timestamp=now - timedelta(days=2, hours=1),
            status=1,
            device_ip="192.168.1.101",
            punch_type="salida",
        ),
        AttendanceRecordDB(
            user_id="3",
            timestamp=now - timedelta(days=1, hours=8),
            status=0,
            device_ip="192.168.1.102",
            punch_type="entrada",
        ),
        AttendanceRecordDB(
            user_id="4",
            timestamp=now - timedelta(hours=4),
            status=2,
            device_ip="192.168.1.101",
            punch_type="entrada_almuerzo",
        ),
        AttendanceRecordDB(
            user_id="5",
            timestamp=now - timedelta(hours=3),
            status=3,
            device_ip="192.168.1.101",
            punch_type="salida_almuerzo",
        ),
    ]


@pytest.fixture
def populated_repo(tmp_path, sample_records):
    """Create a populated buffer repository for testing"""
    db_path = tmp_path / "populated_buffer.db"
    repo = BufferRepository(db_path)
    repo.add_records(sample_records)
    return repo


class TestAttendanceExport:
    """Tests for attendance export functionality"""

    def test_export_csv_streaming(self, app_with_data, sample_records):
        """Test CSV export with streaming"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=csv")

            assert response.status_code == 200
            assert response.mimetype == "text/csv"
            assert "attachment" in response.headers.get("Content-Disposition", "")
            assert "asistencia" in response.headers.get("Content-Disposition", "")
            assert "X-Total-Records" in response.headers
            assert int(response.headers["X-Total-Records"]) == len(sample_records)

            # Parse CSV content (strip BOM if present)
            content = response.data.decode("utf-8")
            if content.startswith("﻿"):
                content = content[1:]
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Check headers
            assert rows[0] == [
                "ID",
                "Usuario",
                "Fecha/Hora",
                "Estado",
                "Dispositivo",
                "Tipo Marca",
                "Sincronizado",
            ]

            # Check data rows (5 data + 1 header = 6 rows)
            assert len(rows) == 6

            # Verify first data row
            assert rows[1][0] != ""  # ID
            assert rows[1][1] in ["1", "2", "3", "4", "5"]  # User ID
            assert rows[1][3] in [
                "Entrada",
                "Salida",
                "Entrada Almuerzo",
                "Salida Almuerzo",
            ]  # Status

    def test_export_csv_with_date_filters(self, app_with_data, sample_records):
        """Test CSV export with date range filters"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            # Filter by start date (today)
            start_date = datetime.now().strftime("%Y-%m-%d")
            response = client.get(f"/api/export/attendance?format=csv&start={start_date}")

            assert response.status_code == 200
            content = response.data.decode("utf-8")
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Should only have records from today (2 records: status 2 and 3 from last few hours)
            assert len(rows) == 3  # header + 2 data rows

    def test_export_csv_with_device_filter(self, app_with_data, sample_records):
        """Test CSV export with device IP filter"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=csv&device_ip=192.168.1.102")

            assert response.status_code == 200
            content = response.data.decode("utf-8")
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Should only have records from device 192.168.1.102 (1 record)
            assert len(rows) == 2  # header + 1 data row
            assert "192.168.1.102" in content

    def test_export_csv_with_user_filter(self, app_with_data, sample_records):
        """Test CSV export with employee/user ID filter"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=csv&user_id=1")

            assert response.status_code == 200
            content = response.data.decode("utf-8")
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Should only have records for user_id=1 (1 record)
            assert len(rows) == 2  # header + 1 data row
            # CSV uses quoted fields, check for "1" in user_id column
            assert '"1"' in content

    def test_export_xlsx_streaming(self, app_with_data, sample_records):
        """Test Excel (xlsx) export with write_only mode"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=xlsx")

            assert response.status_code == 200
            assert (
                response.mimetype
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert "attachment" in response.headers.get("Content-Disposition", "")
            assert "asistencia" in response.headers.get("Content-Disposition", "")
            assert "X-Total-Records" in response.headers
            assert int(response.headers["X-Total-Records"]) == len(sample_records)

            # Parse Excel content
            output = io.BytesIO(response.data)
            wb = load_workbook(output)

            # Check sheet name
            assert "Asistencia" in wb.sheetnames
            ws = wb["Asistencia"]

            # Check headers
            headers = [cell.value for cell in ws[1]]
            assert headers == [
                "ID",
                "Usuario",
                "Fecha/Hora",
                "Estado",
                "Dispositivo",
                "Tipo Marca",
                "Sincronizado",
            ]

            # Check data rows (5 data + 1 header = 6 rows)
            assert ws.max_row == 6

            # Verify first data row
            assert ws.cell(row=2, column=1).value is not None  # ID
            assert ws.cell(row=2, column=2).value in ["1", "2", "3", "4", "5"]  # User ID

    def test_export_xlsx_with_formatting(self, app_with_data, sample_records):
        """Test Excel export has proper formatting (bold headers, colors, borders)"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=xlsx")

            output = io.BytesIO(response.data)
            wb = load_workbook(output)
            ws = wb["Asistencia"]

            # Check header formatting
            header_cell = ws.cell(row=1, column=1)
            assert header_cell.font.bold is True
            # Check fill color (openpyxl uses theme colors, check for blue theme)
            assert header_cell.fill.start_color is not None
            assert header_cell.alignment.horizontal == "center"
            assert header_cell.border.left.style == "thin"

            # Check data cell formatting
            data_cell = ws.cell(row=2, column=1)
            assert data_cell.border.left.style == "thin"

            # Check column widths (openpyxl write_only sets defaults, verify they're set)
            for col in ["A", "B", "C", "D", "E", "F", "G"]:
                assert ws.column_dimensions[col].width is not None
                assert ws.column_dimensions[col].width > 0

    def test_export_invalid_format(self, app_with_data):
        """Test export with invalid format returns 400"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/api/export/attendance?format=pdf")

            assert response.status_code == 400
            assert "Formato no soportado" in response.get_json()["error"]

    def test_export_no_auth(self, client):
        """Test export without authentication returns 401"""
        response = client.get("/api/export/attendance?format=csv")

        assert response.status_code == 401
        assert "No autenticado" in response.get_json()["error"]

    def test_export_no_buffer_repo(self, app):
        """Test export when buffer repo is not available"""
        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            app.buffer_repo = None

            response = client.get("/api/export/attendance?format=csv")

            assert response.status_code == 503
            assert "Buffer repository no disponible" in response.get_json()["error"]

    def test_export_empty_result(self, app, tmp_path):
        """Test export with no matching records"""
        empty_db_path = tmp_path / "empty_buffer.db"
        empty_repo = BufferRepository(empty_db_path)

        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            app.buffer_repo = empty_repo
            response = client.get("/api/export/attendance?format=csv")

            assert response.status_code == 200
            content = response.data.decode("utf-8")
            if content.startswith("﻿"):
                content = content[1:]
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Should only have header row
            assert len(rows) == 1
            assert rows[0] == [
                "ID",
                "Usuario",
                "Fecha/Hora",
                "Estado",
                "Dispositivo",
                "Tipo Marca",
                "Sincronizado",
            ]
            assert response.headers["X-Total-Records"] == "0"

    def test_export_large_dataset_streaming(self, app, tmp_path):
        """Test streaming works for large datasets (no memory blowup)"""
        large_db_path = tmp_path / "large_buffer.db"
        repo = BufferRepository(large_db_path)

        # Add 5000 records (more than chunk size of 1000)
        now = datetime.now()
        records = [
            AttendanceRecordDB(
                user_id=str(i),
                timestamp=now - timedelta(minutes=i),
                status=i % 6,
                device_ip="192.168.1.101",
                punch_type="entrada" if i % 2 == 0 else "salida",
            )
            for i in range(5000)
        ]
        repo.add_records(records)

        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            app.buffer_repo = repo

            # Test CSV streaming
            response = client.get("/api/export/attendance?format=csv")
            assert response.status_code == 200
            assert int(response.headers["X-Total-Records"]) == 5000

            # Verify we can stream the response (check content length)
            content = response.data.decode("utf-8")
            lines = content.strip().split("\n")
            assert len(lines) == 5001  # header + 5000 data rows

    def test_export_attendance_path_alias(self, app_with_data, sample_records):
        """Test /attendance/export path works as alias"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            response = client.get("/attendance/export?format=csv")

            assert response.status_code == 200
            assert response.mimetype == "text/csv"

    def test_export_backward_compat_params(self, app_with_data, sample_records):
        """Test export accepts start_date/end_date params for backward compat"""
        with app_with_data.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            # Use old param names - use wide date range to ensure all records included
            start_date = "2020-01-01"
            end_date = "2030-12-31"
            response = client.get(
                f"/api/export/attendance?format=csv&start_date={start_date}&end_date={end_date}"
            )

            assert response.status_code == 200
            content = response.data.decode("utf-8")
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)

            # Should have all 5 records
            assert len(rows) == 6


class TestAttendanceApi:
    """Tests for attendance API (paginated listing)"""

    def test_api_attendance_pagination(self, app, populated_repo, sample_records):
        """Test paginated attendance API"""
        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            app.buffer_repo = populated_repo

            response = client.get("/api/attendance?page=1&per_page=2")

            assert response.status_code == 200
            data = response.get_json()

            assert "records" in data
            assert "pagination" in data
            assert len(data["records"]) == 2
            assert data["pagination"]["page"] == 1
            assert data["pagination"]["per_page"] == 2
            assert data["pagination"]["total"] == 5
            assert data["pagination"]["total_pages"] == 3

    def test_api_attendance_filters(self, app, populated_repo, sample_records):
        """Test attendance API with filters"""
        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess["user_id"] = "test_user"

            app.buffer_repo = populated_repo

            # Filter by device
            response = client.get("/api/attendance?device_ip=192.168.1.102")
            data = response.get_json()

            assert response.status_code == 200
            assert len(data["records"]) == 1
            assert data["records"][0]["device_ip"] == "192.168.1.102"

            # Filter by user
            response = client.get("/api/attendance?user_id=1")
            data = response.get_json()

            assert len(data["records"]) == 1
            assert data["records"][0]["user_id"] == "1"

    def test_api_attendance_no_auth(self, client):
        """Test attendance API without auth returns 401"""
        response = client.get("/api/attendance")

        assert response.status_code == 401
        assert "No autenticado" in response.get_json()["error"]


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
