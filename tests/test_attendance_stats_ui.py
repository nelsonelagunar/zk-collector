"""Tests para el dashboard de estadisticos /reportes (MYK-242).

El dashboard reutiliza la implementacion de la API de MYK-241, por lo que
los tests clave comparan la respuesta del dashboard contra la de
`/api/attendance/stats` (y su export) para los mismos filtros.
"""

import io
import shutil
import sqlite3
import time
from datetime import date, timedelta
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from src.buffer.repository import BufferRepository
from src.webgui.api_tokens import ApiTokenService
from src.webgui.database import Database
from src.webgui.routes.attendance_stats import attendance_stats_bp
from src.webgui.routes.attendance_stats_ui import attendance_stats_ui_bp

REPO_ROOT = Path(__file__).resolve().parents[1]
MIGRATIONS_DIR = REPO_ROOT / "migrations"
WEBGUI_DIR = REPO_ROOT / "src" / "webgui"
MIGRATION_CHAIN = (
    "001_initial_schema.sql",
    "002_employees.sql",
    "003_medical_leaves.sql",
    "009_api_tokens.sql",
)

AUTH_HEADER: dict[str, str] = {}
START, END = "2026-10-05", "2026-10-09"

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _apply_migrations(conn: sqlite3.Connection) -> None:
    for name in MIGRATION_CHAIN:
        conn.executescript((MIGRATIONS_DIR / name).read_text())


def _build_app(tmp_path: Path) -> Flask:
    shutil.copytree(MIGRATIONS_DIR, tmp_path / "migrations")
    db_path = tmp_path / "data" / "webgui.db"
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    _apply_migrations(conn)
    conn.close()

    app = Flask(
        __name__,
        template_folder=str(WEBGUI_DIR / "templates"),
        static_folder=str(WEBGUI_DIR / "static"),
    )
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test-secret-key"
    app.db = Database(db_path)
    app.buffer_repo = BufferRepository(tmp_path / "data" / "buffer.db")
    app.register_blueprint(attendance_stats_bp, url_prefix="/api/attendance")
    app.register_blueprint(attendance_stats_ui_bp, url_prefix="/reportes")
    return app


@pytest.fixture
def client(tmp_path):
    app = _build_app(tmp_path)
    test_client = app.test_client()

    user_id = app.db.insert(
        "users",
        {"username": "tester", "password_hash": "x", "role": "admin"},
    )
    token = ApiTokenService(app.db).issue_token(
        user_id, "test", ["read", "write"]
    )["token"]
    AUTH_HEADER.clear()
    AUTH_HEADER["Authorization"] = f"Bearer {token}"
    return test_client


@pytest.fixture
def seeded(client):
    db = client.application.db
    buffer_repo = client.application.buffer_repo
    db.insert("devices", {"name": "Reloj 1", "address": "10.0.0.5", "serial_number": "SN1"})
    ana_id = db.insert(
        "employees",
        {
            "employee_code": "E001",
            "first_name": "Ana",
            "last_name": "Lopez",
            "department": "IT",
            "hire_date": "2020-01-01",
        },
    )
    db.insert(
        "employees",
        {
            "employee_code": "E002",
            "first_name": "Bob",
            "last_name": "Suarez",
            "department": "Ventas",
        },
    )
    db.insert(
        "employee_device_links",
        {
            "employee_id": ana_id,
            "device_serial": "SN1",
            "device_user_id": "1",
        },
    )
    with buffer_repo.get_connection() as conn:
        conn.execute(
            "INSERT INTO attendance_records (user_id, timestamp, status, device_ip) "
            "VALUES ('1', '2026-10-05 08:55:00', 0, '10.0.0.5')"
        )
        conn.execute(
            "INSERT INTO attendance_records (user_id, timestamp, status, device_ip) "
            "VALUES ('1', '2026-10-06T09:20:00', 0, '10.0.0.5')"
        )
    return {"ana_id": ana_id}


def _ui_stats(client, query):
    return client.get(f"/reportes/api/stats?{query}")


def _api_stats(client, query):
    return client.get(f"/api/attendance/stats?{query}", headers=dict(AUTH_HEADER))


def _ui_export(client, query):
    return client.get(f"/reportes/export?{query}")


def _api_export(client, query):
    return client.get(
        f"/api/attendance/stats/export?{query}", headers=dict(AUTH_HEADER)
    )


class TestDashboardPage:
    def test_page_renders(self, client, seeded):
        resp = client.get("/reportes/")
        assert resp.status_code == 200
        body = resp.get_data(as_text=True)
        assert "Estadísticos de Asistencia" in body
        # Opciones de los filtros: colaboradores y departamentos activos
        assert "Ana Lopez (E001)" in body
        assert "Bob Suarez (E002)" in body
        assert '<option value="IT">IT</option>' in body
        assert "Ventas" in body

    def test_page_contains_chart_and_filter_markers(self, client, seeded):
        body = client.get("/reportes/").get_data(as_text=True)
        assert "Composición del periodo" in body
        assert "Puntualidad por" in body
        assert "Ausentismo por" in body
        assert "Asistencia por" in body
        markers = ("start_date", "end_date", "employee_id", "department", "group_by")
        for marker in markers:
            assert f'id="{marker}"' in body


class TestStatsJsonParityWithApi:
    def test_ui_json_equals_api_json(self, client, seeded):
        query = f"start_date={START}&end_date={END}"
        ui = _ui_stats(client, query)
        api = _api_stats(client, query)
        assert ui.status_code == 200
        assert api.status_code == 200
        assert ui.get_json() == api.get_json()

    def test_ui_json_equals_api_json_department(self, client, seeded):
        query = f"start_date={START}&end_date={END}&group_by=department"
        assert _ui_stats(client, query).get_json() == _api_stats(client, query).get_json()

    def test_ui_json_equals_api_json_with_filters(self, client, seeded):
        query = f"start_date={START}&end_date={END}&employee_id={seeded['ana_id']}&department=IT"
        assert _ui_stats(client, query).get_json() == _api_stats(client, query).get_json()

    def test_ui_json_default_period_is_current_month(self, client, seeded):
        today = date.today()
        data = _ui_stats(client, "").get_json()
        assert data["period"]["start_date"] == today.replace(day=1).isoformat()
        assert data["period"]["end_date"] == today.isoformat()

    def test_ui_json_requires_no_bearer(self, client, seeded):
        # La UI vive en la sesion WebGUI: sin header Authorization tambien responde
        resp = client.get(f"/reportes/api/stats?start_date={START}&end_date={END}")
        assert resp.status_code == 200

    def test_ui_json_invalid_date_returns_400(self, client):
        resp = _ui_stats(client, "start_date=nope")
        assert resp.status_code == 400
        assert "start_date" in resp.get_json()["error"]

    def test_ui_json_invalid_group_by_returns_400(self, client):
        resp = _ui_stats(client, f"start_date={START}&end_date={END}&group_by=team")
        assert resp.status_code == 400


class TestExport:
    def test_ui_export_csv_identical_to_api(self, client, seeded):
        query = f"start_date={START}&end_date={END}"
        ui = _ui_export(client, query)
        api = _api_export(client, query)
        assert ui.status_code == 200
        assert ui.data == api.data
        assert ui.content_type.startswith(CSV_CONTENT_TYPE)
        assert ui.data.startswith(b"\xef\xbb\xbf")  # BOM utf-8-sig

    def test_ui_export_csv_content(self, client, seeded):
        text = _ui_export(client, f"start_date={START}&end_date={END}").data.decode("utf-8-sig")
        lines = text.strip().splitlines()
        assert lines[0].split(",")[:3] == ["Código", "Colaborador", "Departamento"]
        assert len(lines) == 4  # header + 2 empleados + TOTAL
        assert lines[3].startswith("TOTAL")
        assert "Ana Lopez" in lines[1]

    def test_ui_export_xlsx(self, client, seeded):
        resp = _ui_export(client, f"start_date={START}&end_date={END}&format=xlsx")
        assert resp.status_code == 200
        assert resp.content_type == XLSX_CONTENT_TYPE
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        assert ws.max_row == 4  # header + 2 empleados + TOTAL
        assert ws.cell(row=4, column=1).value == "TOTAL"

    def test_ui_export_same_filters(self, client, seeded):
        resp = _ui_export(
            client, f"start_date={START}&end_date={END}&employee_id={seeded['ana_id']}"
        )
        text = resp.data.decode("utf-8-sig")
        assert "E002" not in text

    def test_ui_export_invalid_params_returns_400(self, client):
        assert _ui_export(client, "start_date=bad").status_code == 400


class TestPerformance:
    """AC MYK-242: el dashboard carga <2s con 10k registros."""

    N_EMPLOYEES = 20
    N_RECORDS = 10_000
    PERF_START, PERF_END = "2026-10-01", "2026-11-30"

    @pytest.fixture
    def seeded_10k(self, client):
        db = client.application.db
        buffer_repo = client.application.buffer_repo
        db.insert("devices", {"name": "Reloj", "address": "10.0.0.5", "serial_number": "SN1"})
        for idx in range(self.N_EMPLOYEES):
            employee_id = db.insert(
                "employees",
                {
                    "employee_code": f"E{idx + 1:03d}",
                    "first_name": f"Nombre{idx + 1}",
                    "last_name": f"Apellido{idx + 1}",
                    "department": f"Depto{idx % 4}",
                    "hire_date": "2020-01-01",
                },
            )
            db.insert(
                "employee_device_links",
                {
                    "employee_id": employee_id,
                    "device_serial": "SN1",
                    "device_user_id": str(idx + 1),
                },
            )

        days = []
        current = date.fromisoformat(self.PERF_START)
        end = date.fromisoformat(self.PERF_END)
        while current <= end:
            if current.weekday() < 5:
                days.append(current.isoformat())
            current += timedelta(days=1)

        records = []
        for i in range(self.N_RECORDS):
            # (user, day) se repite cada lcm(20,43)=860 registros; el minuto y
            # segundo derivados de i rompen la colision del UNIQUE
            # (user_id, timestamp, device_ip) para los 10k registros.
            records.append(
                (
                    str(i % self.N_EMPLOYEES + 1),
                    f"{days[i % len(days)]} 07:{(i // 60) % 60:02d}:{i % 60:02d}",
                    0,
                    "10.0.0.5",
                )
            )
        with buffer_repo.get_connection() as conn:
            conn.executemany(
                "INSERT INTO attendance_records (user_id, timestamp, status, device_ip) "
                "VALUES (?, ?, ?, ?)",
                records,
            )
        return client

    def test_report_under_2s_with_10k_records(self, seeded_10k):
        query = f"start_date={self.PERF_START}&end_date={self.PERF_END}"
        started = time.perf_counter()
        resp = _ui_stats(seeded_10k, query)
        elapsed = time.perf_counter() - started
        assert resp.status_code == 200
        data = resp.get_json()
        assert len(data["rows"]) == self.N_EMPLOYEES
        assert data["totals"]["days_present"] > 0
        assert elapsed < 2.0, f"reporte tardo {elapsed:.2f}s con 10k registros"
