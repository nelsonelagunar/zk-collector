"""Tests para las rutas /api/attendance/stats y /stats/export (MYK-241)."""

import io
import shutil
import sqlite3
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from src.buffer.repository import BufferRepository
from src.webgui.database import Database
from src.webgui.routes.attendance_stats import attendance_stats_bp

REPO_ROOT = Path(__file__).resolve().parents[1]
MIGRATIONS_DIR = REPO_ROOT / "migrations"
MIGRATION_CHAIN = ("001_initial_schema.sql", "002_employees.sql", "003_medical_leaves.sql")

AUTH_HEADER = {"Authorization": "Bearer test-token"}
START, END = "2026-10-05", "2026-10-09"

CSV_CONTENT_TYPE = "text/csv; charset=utf-8"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _apply_migrations(conn: sqlite3.Connection) -> None:
    for name in MIGRATION_CHAIN:
        conn.executescript((MIGRATIONS_DIR / name).read_text())


def _build_app(tmp_path: Path) -> Flask:
    shutil.copytree(MIGRATIONS_DIR, tmp_path / "migrations")
    db_path = tmp_path / "data" / "webgui.db"
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    _apply_migrations(conn)
    conn.close()

    app = Flask(__name__)
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test-secret-key"
    app.db = Database(db_path)
    app.buffer_repo = BufferRepository(tmp_path / "data" / "buffer.db")
    app.register_blueprint(attendance_stats_bp, url_prefix="/api/attendance")
    return app


@pytest.fixture
def client(tmp_path):
    app = _build_app(tmp_path)
    return app.test_client()


@pytest.fixture
def seeded(client):
    db = client.application.db
    buffer_repo = client.application.buffer_repo
    db.insert("devices", {"name": "Reloj 1", "address": "10.0.0.5", "serial_number": "SN1"})
    ana_id = db.insert(
        "employees",
        {
            "employee_code": "E001",
            "first_name": "Ana",
            "last_name": "Lopez",
            "department": "IT",
            "hire_date": "2020-01-01",
        },
    )
    db.insert(
        "employees",
        {
            "employee_code": "E002",
            "first_name": "Bob",
            "last_name": "Suarez",
            "department": "Ventas",
        },
    )
    db.insert(
        "employee_device_links",
        {
            "employee_id": ana_id,
            "device_serial": "SN1",
            "device_user_id": "1",
        },
    )
    with buffer_repo.get_connection() as conn:
        conn.execute(
            "INSERT INTO attendance_records (user_id, timestamp, status, device_ip) "
            "VALUES ('1', '2026-10-05 08:55:00', 0, '10.0.0.5')"
        )
        conn.execute(
            "INSERT INTO attendance_records (user_id, timestamp, status, device_ip) "
            "VALUES ('1', '2026-10-06T09:20:00', 0, '10.0.0.5')"
        )
    return {"ana_id": ana_id}


def _get(client, query, headers=AUTH_HEADER):
    return client.get(f"/api/attendance/stats?{query}", headers=headers)


def _get_export(client, query, headers=AUTH_HEADER):
    return client.get(f"/api/attendance/stats/export?{query}", headers=headers)


class TestAuth:
    def test_requires_authorization_header(self, client):
        resp = client.get("/api/attendance/stats")
        assert resp.status_code == 401

    def test_accepts_bearer_token(self, client, seeded):
        resp = _get(client, f"start_date={START}&end_date={END}")
        assert resp.status_code == 200


class TestStatsReport:
    def test_report_shape(self, client, seeded):
        data = _get(client, f"start_date={START}&end_date={END}").get_json()
        for key in ("period", "filters", "group_by", "work_start", "rows", "totals", "meta"):
            assert key in data
        assert data["period"] == {"start_date": START, "end_date": END}
        assert data["group_by"] == "employee"
        assert data["work_start"] == "09:00:00"
        assert len(data["rows"]) == 2
        ana = data["rows"][0]
        assert ana["employee_code"] == "E001"
        assert ana["days_present"] == 2
        assert ana["late_arrivals"] == 1
        assert ana["punctuality_pct"] == 50.0
        bob = data["rows"][1]
        assert bob["days_present"] == 0
        assert bob["punctuality_pct"] is None

    def test_department_filter_applies(self, client, seeded):
        data = _get(client, f"start_date={START}&end_date={END}&department=IT").get_json()
        assert data["meta"]["employees_total"] == 1
        assert data["rows"][0]["employee_code"] == "E001"

    def test_group_by_department(self, client, seeded):
        data = _get(client, f"start_date={START}&end_date={END}&group_by=department").get_json()
        assert {row["department"] for row in data["rows"]} == {"IT", "Ventas"}

    def test_group_by_company(self, client, seeded):
        data = _get(client, f"start_date={START}&end_date={END}&group_by=company").get_json()
        assert len(data["rows"]) == 1
        assert data["rows"][0]["company"] == "default"

    def test_company_name_from_config(self, client, seeded):
        client.application.config["COMPANY_NAME"] = "ACME"
        data = _get(client, f"start_date={START}&end_date={END}&group_by=company").get_json()
        assert data["rows"][0]["company"] == "ACME"

    def test_work_start_from_config(self, client, seeded):
        client.application.config["ATTENDANCE_WORK_START"] = "08:00"
        data = _get(client, f"start_date={START}&end_date={END}").get_json()
        assert data["work_start"] == "08:00:00"
        ana = data["rows"][0]
        assert ana["late_arrivals"] == 2  # Lun 08:55 y Mar 09:20 tardios con 08:00
        assert ana["punctuality_pct"] == 0.0

    def test_invalid_date_returns_400(self, client):
        resp = _get(client, "start_date=nope")
        assert resp.status_code == 400
        assert "start_date" in resp.get_json()["error"]

    def test_start_after_end_returns_400(self, client):
        resp = _get(client, f"start_date={END}&end_date={START}")
        assert resp.status_code == 400

    def test_invalid_group_by_returns_400(self, client):
        resp = _get(client, f"start_date={START}&end_date={END}&group_by=team")
        assert resp.status_code == 400

    def test_invalid_employee_id_returns_400(self, client):
        resp = _get(client, f"start_date={START}&end_date={END}&employee_id=abc")
        assert resp.status_code == 400

    def test_empty_period_has_zero_rows(self, client):
        data = _get(client, "start_date=2026-01-05&end_date=2026-01-09").get_json()
        assert data["rows"] == []
        assert data["meta"]["employees_total"] == 0


class TestExport:
    def test_export_csv(self, client, seeded):
        resp = _get_export(client, f"start_date={START}&end_date={END}")
        assert resp.status_code == 200
        assert resp.content_type.startswith(CSV_CONTENT_TYPE)
        assert resp.data.startswith(b"\xef\xbb\xbf")  # BOM utf-8-sig
        text = resp.data.decode("utf-8-sig")
        lines = text.strip().splitlines()
        assert lines[0].split(",") == [
            "Código",
            "Colaborador",
            "Departamento",
            "Días Esperados",
            "Días Presentes",
            "Días Ausentes",
            "Días Incapacidad",
            "Días Vacaciones",
            "Retardos",
            "Puntualidad %",
            "Ausentismo %",
        ]
        assert len(lines) == 4  # header + 2 empleados + TOTAL
        assert lines[3].startswith("TOTAL")
        assert "Ana Lopez" in lines[1]

    def test_export_csv_same_filters(self, client, seeded):
        resp = _get_export(
            client, f"start_date={START}&end_date={END}&employee_id={seeded['ana_id']}"
        )
        text = resp.data.decode("utf-8-sig")
        lines = text.strip().splitlines()
        assert len(lines) == 3  # header + Ana + TOTAL
        assert "E002" not in text

    def test_export_xlsx(self, client, seeded):
        resp = _get_export(client, f"start_date={START}&end_date={END}&format=xlsx")
        assert resp.status_code == 200
        assert resp.content_type == XLSX_CONTENT_TYPE
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        assert ws.max_row == 4  # header + 2 + TOTAL
        assert ws.cell(row=1, column=1).value == "Código"
        assert ws.cell(row=2, column=2).value == "Ana Lopez"
        assert ws.cell(row=4, column=1).value == "TOTAL"

    def test_export_group_by_department_csv(self, client, seeded):
        resp = _get_export(client, f"start_date={START}&end_date={END}&group_by=department")
        text = resp.data.decode("utf-8-sig")
        lines = text.strip().splitlines()
        assert lines[0].split(",")[0] == "Departamento"
        assert "IT" in text and "Ventas" in text

    def test_export_rejects_invalid_params(self, client):
        resp = _get_export(client, "start_date=bad")
        assert resp.status_code == 400

    def test_export_requires_auth(self, client):
        resp = client.get("/api/attendance/stats/export")
        assert resp.status_code == 401
